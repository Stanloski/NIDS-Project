import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_defense_document(output_path):
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    PRIMARY = RGBColor(0, 135, 81)     # UNN Green #008751
    SECONDARY = RGBColor(30, 41, 59)   # Slate Navy #1E293B
    TEXT_COLOR = RGBColor(51, 65, 85)  # Neutral Text #334155
    MUTED = RGBColor(100, 116, 139)

    # Style standard paragraph
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = TEXT_COLOR
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(6)

    # Helper for adding styled headers
    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = PRIMARY
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        run = p.add_run(text)
        run.italic = True
        run.font.size = Pt(12)
        run.font.color.rgb = MUTED
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(15)
        run.font.color.rgb = PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12.5)
        run.font.color.rgb = SECONDARY
        return p

    def add_callout(quote_text, speaker="WHAT TO SAY TO THE EXAMINERS:"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F0FDF4") # subtle green tint
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        head = p.add_run(speaker + "\n")
        head.bold = True
        head.font.size = Pt(10)
        head.font.color.rgb = PRIMARY

        body = p.add_run(f'"{quote_text}"')
        body.italic = True
        body.font.size = Pt(11)
        body.font.color.rgb = SECONDARY

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---------------- DOCUMENT CONTENT ----------------

    # Header / Meta
    add_title("NIDS PROJECT DEFENSE MASTER GUIDE")
    add_subtitle("B.Sc Project Defense Cheat Sheet & Oral Examination Preparation Guide\nDepartment of Computer Science • University of Nigeria, Nsukka (UNN)")

    # Metadata Box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Topic:", "Design and Implementation of a Network Intrusion Detection System Using Random Forest Machine Learning Algorithm"),
        ("Candidate / Presenter:", "Okafor Ugochukwu Stanley (Reg No: 2022/249519)"),
        ("Group Members (Group 2):", "Okafor Michael Osinachi, Okafor Ugochukwu Stanley, Okechukwu Daniel Somtochukwu"),
        ("Project Supervisor:", "Dr. U.K Umeh | Academic Year: 2026")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = SECONDARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 1: 2-Minute Opening Pitch
    add_h1("SECTION 1: THE 2-MINUTE OPENING PITCH")
    p = doc.add_paragraph(
        "Deliver this immediately after the panel asks you to introduce yourself and your project. "
        "Keep your back straight, maintain direct eye contact with Dr. Umeh and the examiners, speak audibly, and avoid rushing."
    )
    add_callout(
        "Good morning, distinguished panel of examiners. My name is Okafor Ugochukwu Stanley, presenting on behalf of Group 2. "
        "Our project is titled: Design and Implementation of a Network Intrusion Detection System (NIDS) Using Random Forest Machine Learning Algorithm, "
        "supervised by Dr. U.K. Umeh.\n\n"
        "Traditional perimeter firewalls rely on static signatures and static port rules. The moment an attacker modifies packet payloads or executes a zero-day exploit, rule-based defenses go blind. "
        "To solve this, we engineered an intelligent, anomaly-based NIDS powered by an ensemble Random Forest classifier.\n\n"
        "We trained our system on the benchmark NSL-KDD dataset across 41 network connection attributes, capturing protocol behaviors, byte exchange rates, and host-level traffic patterns. "
        "We then operationalized the machine learning pipeline into a web security operations center—RADAR WATCHTOWER—featuring live stream simulation, single-connection diagnosis, automated batch CSV classification, and tamper-resistant audit logging.\n\n"
        "With your permission, I would like to demonstrate how the system identifies and isolates network intrusions in real time."
    )

    # Section 2: Key Numbers & Model Performance
    add_h1("SECTION 2: KEY NUMBERS & METRICS (MEMORIZE THESE)")
    p = doc.add_paragraph(
        "Examiners frequently test whether you know your exact experimental results or simply copied a template. "
        "These are the exact verified figures from your training run on the official NSL-KDD benchmark:"
    )

    metrics_table = doc.add_table(rows=9, cols=3)
    metrics_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_headers = ["Parameter / Metric", "Exact Value", "Defense Explanation / Significance"]
    for j, h in enumerate(table_headers):
        cell = metrics_table.cell(0, j)
        set_cell_background(cell, "008751")
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    metrics_rows = [
        ("Dataset Used", "NSL-KDD Benchmark", "125,973 training records (KDDTrain+), 22,544 test records (KDDTest+)"),
        ("Model & Configuration", "Random Forest (n=100)", "Ensemble of 100 Decision Trees, ColumnTransformer, OrdinalEncoder"),
        ("Test Precision", "96.68%", "Crucial: When the alarm sounds, it is 96.7% accurate. Negligible false alarms."),
        ("False Positive Rate", "2.82%", "Only 274 normal connections out of 9,711 were misclassified."),
        ("Test Accuracy", "77.28%", "Evaluated on KDDTest+, which includes 17 novel zero-day attacks."),
        ("Test Recall", "62.23%", "Coverage across both known training attacks and brand new test variants."),
        ("F1-Score", "75.72%", "Harmonic mean of precision and recall on the hardest benchmark test set."),
        ("Top Discriminant Features", "src_bytes (18.5%), flag (13.1%), dst_bytes (9.9%)", "Connection payload bytes and TCP flags are the highest predictors of attack.")
    ]

    for i, (col1, col2, col3) in enumerate(metrics_rows):
        row = metrics_table.rows[i + 1]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width = Inches(1.8)
        c1.width = Inches(1.5)
        c2.width = Inches(3.2)
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for c in (c0, c1, c2):
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 100, 100)

        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(col1); r0.bold = True; r0.font.size = Pt(9.5)

        p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(col2); r1.bold = True; r1.font.size = Pt(9.5); r1.font.color.rgb = PRIMARY

        p2 = c2.paragraphs[0]; p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(col3); r2.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 3: Live System Demo Walkthrough
    add_h1("SECTION 3: STEP-BY-STEP LIVE DEMO SCRIPT")
    p = doc.add_paragraph(
        "Follow this exact script when demonstrating your system on screen. "
        "Narrate your actions clearly before clicking."
    )

    steps = [
        ("Step 1: Security & Authentication", 
         "Show the login page. Mention: 'We implemented OWASP authentication best practices. Passwords use PBKDF2 hashing with SHA-256 and salt. After 5 consecutive failed attempts, the system enforces a 15-minute brute-force lockout.' Log in with your admin credentials."),
        
        ("Step 2: RADAR Operations Dashboard",
         "The dashboard loads with UNN aesthetics (Dark/Light mode). Point to the KPI cards: 'The dashboard gives the network administrator real-time situational awareness: total monitored packets, attack vs normal distribution, and live Chart.js threat curves. The system runs an asynchronous traffic simulator simulating real-world gateway throughput.'"),
        
        ("Step 3: Single Packet Live Diagnosis (Show-Stopper)",
         "Navigate to Manual Predict. Click the preset 'Normal HTTP Traffic'. Say: 'Here, standard TCP 3-way handshake with SF flag and high destination bytes results in a clean NORMAL verdict.' Next, click the preset 'DoS SYN Flood Attack'. Say: 'Notice that destination bytes drop to 0, status flag becomes S0 (SYN sent, no ACK), and same-host connection rates spike. The Random Forest instantly flags this as an INTRUSION with 99% probability.'"),
         
        ("Step 4: Batch CSV Packet Ingestion",
         "Navigate to Batch Upload. Drag and drop sample_upload.csv. Say: 'In enterprise deployments, gigabytes of network logs are ingested asynchronously. Our batch pipeline vectorizes 41 features per row, executes high-speed inference, assigns threat risk levels, and updates the database.'"),
         
        ("Step 5: Forensic Audit Logs & Export",
         "Navigate to History. Show the filterable table. Say: 'Every verdict is timestamped and logged with its complete feature vector for forensic analysis. The administrator can filter by classification status or export the audit trail directly to CSV.'")
    ]

    for title, desc in steps:
        add_h2(title)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(desc)
        r.font.size = Pt(10.5)

    # Section 4: The 8 Tough Defense Questions & Winning Answers
    add_h1("SECTION 4: THE 8 TOUGH QUESTIONS EXAMINERS LOVE TO ASK")

    qa_list = [
        ("Question 1: Why did you choose Random Forest over Deep Learning (CNN/LSTM) or Logistic Regression?",
         "Sir/Ma, we evaluated that trade-off carefully: "
         "1) Logistic Regression is a linear model that fails to capture complex non-linear interactions across 41 network attributes. "
         "2) Deep Learning models (like CNN/LSTM) act as black boxes, require GPU compute, and exhibit inference latency too high for packet inspection at line speed. "
         "3) Random Forest is an ensemble of decision trees using bagging (bootstrap aggregation). It natively handles mixed continuous and categorical data, resists overfitting, provides transparent feature importance (Explainable AI), and executes predictions in under 2 milliseconds on commodity CPU hardware."),
        
        ("Question 2: Isn't 77% accuracy too low for machine learning?",
         "That is an insightful question, Sir. In academic literature, papers reporting 98% or 99% accuracy on KDD typically perform a random split on KDDTrain+, which creates severe data leakage. "
         "Our research evaluated the model against the official, standardized KDDTest+ benchmark. KDDTest+ was intentionally designed by cybersecurity researchers to contain 17 novel zero-day attack variants (such as mscan, udpstorm, and ps) that never appeared in the training set. "
         "In network security, Precision (96.68%) and False Positive Rate (2.82%) are vastly more critical than raw accuracy. Our system ensures legitimate network traffic is virtually never blocked, and any triggered alert represents a high-probability security incident."),
         
        ("Question 3: How is this NIDS different from a standard Firewall?",
         "A Firewall acts as a security guard checking static ID badges at the door—filtering based on static IP addresses, port numbers, or protocol types. If malicious traffic is disguised over port 80 (HTTP) or port 443 (HTTPS), the firewall allows it through. "
         "Our NIDS functions like a security camera with behavioral analysis inside the building. It inspects packet dynamics, payload byte ratios, error rates, and connection frequency to detect malicious behavior regardless of which port or address the traffic originates from."),
         
        ("Question 4: How does your pipeline handle categorical text features like protocol_type or service?",
         "We built a Scikit-Learn Pipeline integrating a ColumnTransformer. Categorical features—protocol_type (tcp/udp/icmp), service, and flag—are encoded using an OrdinalEncoder with handle_unknown='use_encoded_value' set to -1. "
         "This guarantees that if the system encounters an unmapped service during live inspection, the application will not crash; it assigns a fallback token and completes classification safely."),
         
        ("Question 5: What are the most important features that identify an intrusion?",
         "According to Gini impurity feature importance derived from our 100 decision trees: "
         "1) src_bytes (18.53%): Source-to-destination payload bytes. Abnormal surges or zero-byte anomalies are typical of DoS and scanning attacks. "
         "2) flag (13.14%): TCP connection status flags. Normal traffic establishes connection (SF), whereas SYN floods or port scans leave connections half-open (S0 or REJ). "
         "3) dst_bytes (9.91%): Destination-to-source bytes. "
         "Together, these top features drive over 41% of the model's predictive capability."),
         
        ("Question 6: What happens if the database locks or network connection drops?",
         "Our system decouples model inference from database persistence. The Random Forest model (rf_model.joblib) is deserialized into server RAM upon application startup. Real-time classification executes completely in-memory. "
         "Furthermore, all web assets—fonts, icons, and Chart.js—are bundled locally. The application can run entirely offline in an air-gapped security operations center without external internet dependencies."),
         
        ("Question 7: How did you test and validate the system before deployment?",
         "We authored an automated test suite containing 17 comprehensive test cases (test_suite.py) spanning unit testing, database CRUD integrity, authentication lockouts, single and batch ML inference, and edge-case payload validation. All 17 tests execute with a 100% pass rate."),
         
        ("Question 8: What was YOUR specific personal contribution to this project?",
         "As a core member of Group 2, my contributions centered on: "
         "1) Constructing the machine learning preprocessing and training pipeline on the NSL-KDD dataset. "
         "2) Developing the Flask backend architecture, database schema, and OWASP security controls. "
         "3) Engineering the operations dashboard interface and automating the 17-part test suite to verify end-to-end reliability.")
    ]

    for q, a in qa_list:
        add_h2(q)
        add_callout(a, speaker="RECOMMENDED VERBATIM RESPONSE:")

    # Section 5: Pre-Flight Defense Checklist
    add_h1("SECTION 5: PRE-FLIGHT DEFENSE DAY CHECKLIST")
    checklist_items = [
        "Laptop fully charged (100%) with power adapter plugged in.",
        "Flask app tested and running: open terminal, run 'python app.py', test http://127.0.0.1:5000.",
        "Browser opened in clean full-screen mode (F11), bookmarks bar hidden.",
        "GitHub repository (https://github.com/Stanloski/NIDS-Project) open in a separate tab to show version control if asked.",
        "Print out or review this guide 30 minutes before entering the examination hall.",
        "Rule of Delivery: Speak calmly, avoid interrupting examiners, acknowledge their questions with 'Thank you, Sir/Ma', and answer directly."
    ]
    for item in checklist_items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        r.font.size = Pt(10.5)

    doc.save(str(output_path))
    print(f"Successfully generated defense guide at: {output_path}")

if __name__ == "__main__":
    out = Path(r"c:\Users\Dell\Desktop\NIDS Project\NIDS_Project_Defense_Master_Guide.docx")
    create_defense_document(out)
